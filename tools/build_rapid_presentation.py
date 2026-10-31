"""Build an editable rapid presentation from a specimen PowerPoint template.

The template retains one specimen slide for every supported layout. This script
opens the template, removes those specimen slides, adds product slides using
named layouts across all slide masters, populates placeholders, adds speaker
notes and source traceability, and saves a new PPTX.

A slide names either a template "layout", whose placeholders it fills, or a
"pattern" from rapid_patterns.py, which draws an expressive slide in the deck's
color "family" (Blue or Green).
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from rapid_patterns import PatternError, add_pattern_slide


TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def layout_by_name(presentation: Presentation, name: str):
    matches = [
        layout
        for master in presentation.slide_masters
        for layout in master.slide_layouts
        if layout.name == name
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected exactly one layout named {name!r}; found {len(matches)}"
        )
    return matches[0]


def remove_all_slides(presentation: Presentation) -> None:
    """Remove specimen slides while retaining their masters and layouts."""
    slide_id_list = presentation.slides._sldIdLst
    for slide_id in list(slide_id_list):
        relationship_id = slide_id.rId
        presentation.part.drop_rel(relationship_id)
        slide_id_list.remove(slide_id)


def title_placeholder(slide):
    for shape in slide.placeholders:
        if shape.placeholder_format.type in TITLE_TYPES:
            return shape
    raise ValueError(f"Layout {slide.slide_layout.name!r} has no title placeholder")


def subtitle_placeholder(slide):
    """Return the subtitle, or the first body placeholder on layouts such as
    Section that use a body placeholder for their subtitle."""
    for shape in slide.placeholders:
        if shape.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
            return shape
    content = content_placeholders(slide)
    if content:
        return content[0]
    raise ValueError(f"Layout {slide.slide_layout.name!r} has no subtitle placeholder")


def content_placeholders(slide):
    placeholders = [
        shape
        for shape in slide.placeholders
        if shape.placeholder_format.type not in TITLE_TYPES
        and shape.placeholder_format.type != PP_PLACEHOLDER.SUBTITLE
        and shape.placeholder_format.type
        not in {
            PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.FOOTER,
            PP_PLACEHOLDER.SLIDE_NUMBER,
        }
    ]
    return sorted(placeholders, key=lambda shape: shape.placeholder_format.idx)


def set_text(shape, blocks: list[dict[str, Any]]) -> None:
    text_frame = shape.text_frame
    text_frame.clear()
    for index, block in enumerate(blocks):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = block["text"]
        paragraph.level = int(block.get("level", 0))
        if block.get("bold"):
            for run in paragraph.runs:
                run.font.bold = True


def set_notes(slide, speaker_notes: str, sources: list[str], core_release: str) -> None:
    notes = speaker_notes.strip()
    traceability = "\n\nSources:\n" + "\n".join(f"- {source}" for source in sources)
    release = f"\n\nCore release: {core_release}"
    slide.notes_slide.notes_text_frame.text = notes + traceability + release


def add_slide(
    presentation: Presentation,
    slide_spec: dict[str, Any],
    core_release: str,
    family: str | None = None,
    base: Path = Path("."),
):
    if "pattern" in slide_spec:
        chosen = slide_spec.get("family", family)
        if not chosen:
            raise ValueError(
                f"Pattern slide {slide_spec['pattern']!r} needs a 'family', on the slide or the deck"
            )
        slide = add_pattern_slide(presentation, slide_spec, chosen, base, layout_by_name)
        set_notes(slide, slide_spec["speaker_notes"], slide_spec["sources"], core_release)
        return slide

    layout = layout_by_name(presentation, slide_spec["layout"])
    slide = presentation.slides.add_slide(layout)
    if slide_spec.get("title"):
        title_placeholder(slide).text = slide_spec["title"]

    available_content = content_placeholders(slide)
    if "subtitle" in slide_spec:
        subtitle = subtitle_placeholder(slide)
        subtitle.text = slide_spec["subtitle"]
        subtitle_index = subtitle.placeholder_format.idx
        available_content = [
            shape
            for shape in available_content
            if shape.placeholder_format.idx != subtitle_index
        ]

    requested_content = slide_spec.get("content", [])
    if len(requested_content) != len(available_content):
        raise ValueError(
            f"Slide {slide_spec.get('title', '(untitled)')!r} requests {len(requested_content)} content "
            f"regions but layout {slide_spec['layout']!r} has {len(available_content)}"
        )
    for shape, blocks in zip(available_content, requested_content):
        set_text(shape, blocks)

    set_notes(
        slide,
        slide_spec["speaker_notes"],
        slide_spec["sources"],
        core_release,
    )
    return slide


def main() -> None:
    args = parse_args()
    source = json.loads(args.source.read_text(encoding="utf-8-sig"))
    presentation = Presentation(str(args.template))
    remove_all_slides(presentation)

    family = source.get("family")
    base = args.source.resolve().parent
    for number, slide_spec in enumerate(source["slides"], start=1):
        try:
            add_slide(presentation, slide_spec, source["core_release"], family, base)
        except PatternError as error:
            raise SystemExit(f"Slide {number}: {error}") from None

    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(args.output))
    print(f"Created {args.output} with {len(presentation.slides)} slides")


if __name__ == "__main__":
    main()

